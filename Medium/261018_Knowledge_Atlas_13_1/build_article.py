"""Build the editable and plain-text versions from the illustrated article.

Run with the Codex primary Python runtime (python-docx and lxml required).
The five JPEG screenshots are original live-browser captures, never retouched.
"""
from pathlib import Path
import hashlib
import json
import re
from lxml import html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Medium_Article_Atlas_13_1.html"
tree = html.fromstring(SOURCE.read_text(encoding="utf-8"))
article = tree.xpath("//article")[0]

doc = Document()
# Strip Word template title borders before adding the article.
for border in list(doc.styles.element.xpath(".//w:pBdr")):
    border.getparent().remove(border)
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.65)
section.bottom_margin = Inches(.65)
section.left_margin = Inches(.7)
section.right_margin = Inches(.7)

for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption", "List Number"]:
    style = doc.styles[name]
    style.font.name = "Georgia" if name in {"Normal", "List Number"} else "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    if style.element.rPr is not None:
        for color in style.element.rPr.findall(qn("w:color")):
            for attr in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn("w:" + attr), None)
    style.paragraph_format.widow_control = True

normal = doc.styles["Normal"]
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(7)
doc.styles["Title"].font.size = Pt(27)
doc.styles["Title"].font.bold = True
doc.styles["Title"].paragraph_format.space_after = Pt(12)
doc.styles["Subtitle"].font.size = Pt(13)
doc.styles["Subtitle"].paragraph_format.space_after = Pt(9)
for name in ("Heading 1", "Heading 2"):
    doc.styles[name].font.size = Pt(17)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(17)
    doc.styles[name].paragraph_format.space_after = Pt(9)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles["Caption"].font.size = Pt(9.5)
doc.styles["Caption"].font.bold = False
doc.styles["Caption"].font.italic = False
doc.styles["Caption"].paragraph_format.line_spacing = 1.1
doc.styles["Caption"].paragraph_format.space_after = Pt(12)
doc.styles["List Number"].font.size = Pt(11)
doc.styles["List Number"].paragraph_format.line_spacing = 1.12
doc.styles["List Number"].paragraph_format.space_after = Pt(7)

doc.core_properties.title = article.find("h1").text_content()
doc.core_properties.author = "Enrico Papalini"
doc.core_properties.subject = "Medium article about Knowledge Atlas 13.1 and Non-Deterministic Software Engineering, Second Edition"
doc.core_properties.keywords = "Knowledge Atlas, software engineering, AI coding, knowledge graph, NDSE2"
doc.core_properties.comments = ""

def append_text(paragraph, text, bold=False, italic=False, link=None):
    if not text:
        return
    if link:
        rel_id = paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), rel_id)
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "20576A")
        rpr.append(color)
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rpr.append(underline)
        if bold:
            rpr.append(OxmlElement("w:b"))
        if italic:
            rpr.append(OxmlElement("w:i"))
        r.append(rpr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        hyperlink.append(r)
        paragraph._p.append(hyperlink)
    else:
        run = paragraph.add_run(text)
        run.bold = True if bold else None
        run.italic = True if italic else None

def render_inline(element, paragraph, bold=False, italic=False, link=None):
    bold = bold or element.tag in ("strong", "b")
    italic = italic or element.tag in ("em", "i")
    link = element.get("href") if element.tag == "a" else link
    append_text(paragraph, element.text, bold, italic, link)
    for child in element:
        render_inline(child, paragraph, bold, italic, link)
        append_text(paragraph, child.tail, bold, italic, link)

text_blocks = []
figure_data = []
links = []
for el in article:
    tag = el.tag
    if tag == "figure":
        im = el.xpath(".//img")[0]
        path = ROOT / im.get("src")
        assert path.is_file() and path.stat().st_size > 0, path
        caption = el.find("figcaption").text_content().strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        shape = p.add_run().add_picture(str(path), width=Inches(7.1))
        shape._inline.docPr.set("descr", im.get("alt"))
        shape._inline.docPr.set("title", f"Knowledge Atlas figure {el.get('data-figure')}")
        p = doc.add_paragraph(caption, style="Caption")
        text_blocks.append(f"[IMAGE {el.get('data-figure')}: {path.name}]\nCaption: {caption}\nAlt text: {im.get('alt')}")
        figure_data.append({"figure": int(el.get("data-figure")), "file":im.get("src"), "caption":caption, "alt_text":im.get("alt"), "sha256":hashlib.sha256(path.read_bytes()).hexdigest()})
    elif tag == "ol":
        numbered = []
        for i, li in enumerate(el.findall("li"), 1):
            p = doc.add_paragraph(style="List Number")
            render_inline(li, p)
            numbered.append(f"{i}. {li.text_content().strip()}")
        text_blocks.append("\n\n".join(numbered))
    elif tag in ("h1", "h2", "p"):
        css = el.get("class", "")
        style = {"h1":"Title", "h2":"Heading 1"}.get(tag, "Normal")
        if css == "subtitle":
            style = "Subtitle"
        elif css == "credit":
            style = "Caption"
        p = doc.add_paragraph(style=style)
        if css == "byline":
            p.paragraph_format.space_after = Pt(15)
        if tag == "h1":
            p.add_run(el.text_content().replace("Writing About", "Writing\nAbout"))
        else:
            render_inline(el, p)
        content = el.text_content().strip()
        refs = []
        for a in el.xpath(".//a"):
            href = a.get("href")
            if href not in [x[1] for x in links]:
                links.append((a.text_content().strip(), href))
            index = next(i+1 for i,x in enumerate(links) if x[1] == href)
            refs.append(str(index))
        if refs:
            content += " [" + ", ".join(refs) + "]"
        text_blocks.append(content)

doc.save(ROOT / "Medium_Article_Atlas_13_1.docx")
text_blocks.append("ARTICLE LINKS\n\n" + "\n\n".join(f"[{i}] {label}\n{url}" for i, (label, url) in enumerate(links,1)))
(ROOT / "Medium_Article_Atlas_13_1.txt").write_text("\n\n".join(text_blocks)+"\n", encoding="utf-8")
(ROOT / "Screenshot_Captions_and_Alt_Text.json").write_text(json.dumps(figure_data, ensure_ascii=False, indent=2)+"\n",encoding="utf-8")
print(json.dumps({"article_words":len(re.findall(r"\S+",article.text_content())),"figures":len(figure_data),"links":len(links),"docx":str(ROOT / "Medium_Article_Atlas_13_1.docx")},indent=2))
